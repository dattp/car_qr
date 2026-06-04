"""Core logic: đọc Excel, sinh HTML, sinh PDF QR codes."""

from __future__ import annotations

import html
import io
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Optional

import openpyxl
import qrcode
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

PREFERRED_SHEET = "Lịch bảo dưỡng (2)"
HEADER_ROWS = 2
VIN_COL_IDX = 2

LABELS = [
    None, "STT", "VIN", "Số máy", "Xe", "Phiên bản", "Màu",
    "Ngày nhập kho", "Ngày lưu kho (ngày)", "Ngày bảo dưỡng",
    "Tuần", "Mốc KT",
    "Ngày bảo dưỡng 3 tháng", "Ngày bảo dưỡng 6 tháng",
    "Ngày bảo dưỡng 9 tháng", "Ngày bảo dưỡng 12 tháng", "Note",
]

ProgressFn = Callable[[int, int, str], None]


def _fmt(v) -> str:
    if v is None or v == "":
        return "-"
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v)


def _pick_sheet(wb):
    if PREFERRED_SHEET in wb.sheetnames:
        return wb[PREFERRED_SHEET]
    return wb[wb.sheetnames[0]]


def load_vehicles(xlsx_path: Path) -> list[tuple]:
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True, read_only=True)
    ws = _pick_sheet(wb)

    out: list[tuple] = []
    seen: set[str] = set()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= HEADER_ROWS:
            continue
        if len(row) <= VIN_COL_IDX:
            continue
        vin = row[VIN_COL_IDX]
        if not vin or not isinstance(vin, str):
            continue
        vin = vin.strip()
        if len(vin) < 5 or vin in seen:
            continue
        seen.add(vin)
        out.append(row)
    return out


def _html_for(row: tuple, updated: str) -> str:
    vin = row[VIN_COL_IDX].strip()
    model = html.escape((row[4] or "").strip() if row[4] else "")
    version = html.escape((row[5] or "").strip() if row[5] else "")
    title = f"{model} {version} - {vin}".strip(" -")

    rows_html = []
    for i, label in enumerate(LABELS):
        if label is None:
            continue
        value = html.escape(_fmt(row[i]) if i < len(row) else "-")
        rows_html.append(f"<tr><th>{html.escape(label)}</th><td>{value}</td></tr>")
    table = "\n      ".join(rows_html)

    return f"""<!DOCTYPE html>
<html lang="vi">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
            margin: 0; padding: 16px; background: #f4f6f8; color: #1a1a1a; }}
    .card {{ max-width: 600px; margin: 0 auto; background: #fff; border-radius: 12px;
             padding: 20px; box-shadow: 0 2px 12px rgba(0,0,0,.06); }}
    h1 {{ font-size: 1.15rem; margin: 0 0 4px; color: #003a70; }}
    .vin {{ font-family: ui-monospace, "SF Mono", Menlo, monospace; font-size: .9rem;
            color: #555; margin-bottom: 16px; word-break: break-all; }}
    table {{ width: 100%; border-collapse: collapse; }}
    th, td {{ padding: 10px 8px; text-align: left; border-bottom: 1px solid #eee;
             font-size: .95rem; vertical-align: top; }}
    th {{ font-weight: 500; color: #555; width: 48%; }}
    td {{ font-weight: 600; word-break: break-word; }}
    tr:last-child th, tr:last-child td {{ border-bottom: none; }}
    footer {{ text-align: center; margin-top: 16px; font-size: .72rem; color: #999; }}
  </style>
</head>
<body>
  <div class="card">
    <h1>Thông tin xe</h1>
    <div class="vin">VIN: {html.escape(vin)}</div>
    <table>
      {table}
    </table>
  </div>
  <footer>Cập nhật: {html.escape(updated)}</footer>
</body>
</html>
"""


def _index_html(vehicles: list[tuple]) -> str:
    items = "\n      ".join(
        f'<li><a href="{html.escape(r[VIN_COL_IDX].strip())}.html">'
        f'{html.escape(r[VIN_COL_IDX].strip())} '
        f'— {html.escape(r[4] or "")} {html.escape(r[5] or "")}</a></li>'
        for r in vehicles
    )
    return f"""<!DOCTYPE html>
<html lang="vi"><head><meta charset="utf-8">
<title>Danh sách xe</title>
<style>body{{font-family:sans-serif;max-width:800px;margin:20px auto;padding:0 16px}}
li{{margin:4px 0}}</style></head>
<body><h1>Danh sách xe ({len(vehicles)})</h1><ul>
      {items}
</ul></body></html>
"""


def write_html(
    vehicles: list[tuple],
    dist_dir: Path,
    progress: Optional[ProgressFn] = None,
) -> None:
    dist_dir.mkdir(parents=True, exist_ok=True)
    updated = datetime.now().strftime("%d/%m/%Y")
    total = len(vehicles)
    for idx, row in enumerate(vehicles, 1):
        vin = row[VIN_COL_IDX].strip()
        (dist_dir / f"{vin}.html").write_text(_html_for(row, updated), encoding="utf-8")
        if progress and (idx % 50 == 0 or idx == total):
            progress(idx, total, f"Đang tạo HTML: {idx}/{total}")
    (dist_dir / "index.html").write_text(_index_html(vehicles), encoding="utf-8")


def write_pdf(
    vehicles: list[tuple],
    pdf_path: Path,
    base_url: str,
    progress: Optional[ProgressFn] = None,
) -> None:
    cols, rows = 3, 4
    per_page = cols * rows
    margin = 1.0 * cm
    label_h = 0.55 * cm
    gap = 0.25 * cm

    page_w, page_h = A4
    cell_w = (page_w - 2 * margin) / cols
    cell_h = (page_h - 2 * margin) / rows
    qr_size = min(cell_w - gap, cell_h - label_h - gap)

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    c.setTitle("QR codes - Xe")

    total = len(vehicles)
    base = base_url.rstrip("/")
    for idx, vrow in enumerate(vehicles):
        vin = vrow[VIN_COL_IDX].strip()

        page_idx = idx % per_page
        if idx > 0 and page_idx == 0:
            c.showPage()

        col = page_idx % cols
        row_in = page_idx // cols

        qr = qrcode.QRCode(
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=10, border=1,
        )
        qr.add_data(f"{base}/{vin}.html")
        qr.make(fit=True)
        img = qr.make_image()
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)

        cell_x = margin + col * cell_w
        cell_top = page_h - margin - row_in * cell_h
        qr_x = cell_x + (cell_w - qr_size) / 2
        qr_y = cell_top - qr_size - gap / 2

        c.drawImage(ImageReader(buf), qr_x, qr_y, qr_size, qr_size)
        c.setFont("Helvetica", 7)
        c.drawCentredString(cell_x + cell_w / 2, qr_y - label_h + 0.15 * cm, vin)

        if progress and ((idx + 1) % 50 == 0 or (idx + 1) == total):
            progress(idx + 1, total, f"Đang tạo QR PDF: {idx + 1}/{total}")

    c.save()


def generate(
    xlsx_path: Path,
    output_dir: Path,
    subdomain: str,
    progress: Optional[ProgressFn] = None,
) -> tuple[Path, Path, int]:
    """Sinh dist/ + qr_codes.pdf. Trả về (dist_dir, pdf_path, vehicle_count)."""
    subdomain = subdomain.strip().strip("/")
    if not subdomain:
        raise ValueError("Cần nhập tên subdomain Netlify")
    base_url = f"https://{subdomain}.netlify.app"

    if progress:
        progress(0, 1, "Đang đọc file Excel...")

    vehicles = load_vehicles(xlsx_path)
    if not vehicles:
        raise ValueError("Không tìm thấy xe nào có VIN hợp lệ trong file")

    dist_dir = output_dir / "dist"
    pdf_path = output_dir / "qr_codes.pdf"

    write_html(vehicles, dist_dir, progress)
    write_pdf(vehicles, pdf_path, base_url, progress)

    if progress:
        progress(len(vehicles), len(vehicles), f"Hoàn thành! {len(vehicles)} xe.")

    return dist_dir, pdf_path, len(vehicles)
